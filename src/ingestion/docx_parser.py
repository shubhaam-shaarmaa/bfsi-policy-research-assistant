"""DOCX Document Parser using python-docx.

Parses banking SOPs, policy Word documents, and guidelines.
Extracts headings, paragraphs, and tables with metadata.
"""

from pathlib import Path
import docx
from src.core.logging import setup_logger
from src.core.models import (
    DocumentMetadata,
    DocumentType,
    ExtractedSection,
    ParsedDocument,
)
from src.ingestion.parser_base import BaseDocumentParser

logger = setup_logger("docx_parser")


class DocxDocumentParser(BaseDocumentParser):
    """Parser for Microsoft Word (.docx) documents."""

    def supports(self, file_extension: str) -> bool:
        return file_extension.lower() in [".docx"]

    def parse(self, file_path: Path) -> ParsedDocument:
        if not file_path.exists():
            raise FileNotFoundError(f"DOCX file not found: {file_path}")

        file_hash = self.compute_file_hash(file_path)
        doc_id = f"doc_{file_hash[:12]}"

        try:
            doc = docx.Document(file_path)
        except Exception as e:
            raise ValueError(f"Failed to read DOCX file {file_path}: {e}") from e

        # Extract title from core properties or filename
        title = doc.core_properties.title or file_path.stem.replace("_", " ").title()

        sections: list[ExtractedSection] = []
        full_text_parts: list[str] = []

        current_heading = "Overview"
        current_content: list[str] = []
        section_counter = 1

        # Estimated page approximation for DOCX (~450 words per page)
        word_count_accum = 0

        for para in doc.paragraphs:
            text = para.text.strip()
            if not text:
                continue

            full_text_parts.append(text)
            words_in_para = len(text.split())
            word_count_accum += words_in_para
            current_page = max(1, (word_count_accum // 450) + 1)

            # Check if paragraph has Heading style
            style_name = para.style.name if para.style else ""
            if "Heading" in style_name or (len(text) < 70 and text.isupper()):
                if current_content:
                    sections.append(
                        ExtractedSection(
                            section_id=f"{doc_id}#s{section_counter:03d}",
                            heading=current_heading,
                            page_number=current_page,
                            content="\n".join(current_content).strip(),
                        )
                    )
                    section_counter += 1
                    current_content = []
                current_heading = text
            else:
                current_content.append(text)

        # Also extract text from tables
        for table in doc.tables:
            table_rows: list[str] = []
            for row in table.rows:
                row_cells = [cell.text.strip() for cell in row.cells]
                table_rows.append(" | ".join(row_cells))
            if table_rows:
                table_text = "[TABLE]\n" + "\n".join(table_rows) + "\n[/TABLE]"
                current_content.append(table_text)
                full_text_parts.append(table_text)

        if current_content:
            current_page = max(1, (word_count_accum // 450) + 1)
            sections.append(
                ExtractedSection(
                    section_id=f"{doc_id}#s{section_counter:03d}",
                    heading=current_heading,
                    page_number=current_page,
                    content="\n".join(current_content).strip(),
                )
            )

        total_pages = max(1, (word_count_accum // 450) + 1)

        metadata = DocumentMetadata(
            doc_id=doc_id,
            title=title,
            issuer="Banking Organization / Policy Unit",
            circular_number=None,
            date=None,
            doc_type=DocumentType.SOP,
            source_filename=file_path.name,
            source_path=str(file_path),
            file_hash=file_hash,
            file_format="docx",
            total_pages=total_pages,
        )

        return ParsedDocument(
            metadata=metadata,
            sections=sections,
            raw_text="\n\n".join(full_text_parts),
            page_count=total_pages,
            parsing_warnings=[],
        )
